"""Pruebas funcionales de listados, búsqueda, exportación y consulta.

Cubren el alcance por delegación en el servidor, la búsqueda parcial, la
exportación que respeta el filtro y la ficha de consulta.
"""

from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.ciudadanos.models import Vecino

pytestmark = pytest.mark.django_db


def _vecino(nombre, rut, territorio):
    return Vecino.objects.create(nombre=nombre, rut=rut, territorio=territorio)


def test_el_funcionario_solo_ve_su_delegacion(cliente, crear_usuario, centro, mar):
    _vecino("Ana Centro", "11111111-1", centro)
    _vecino("Beto Mar", "12222222-5", mar)
    usuario = crear_usuario(
        "funcionario@test.cl", codigo="funcionario", delegacion=centro
    )
    cliente.force_login(usuario)

    respuesta = cliente.get(reverse("ciudadanos:vecinos"))

    assert respuesta.status_code == 200
    assert [v.nombre for v in respuesta.context["object_list"]] == ["Ana Centro"]


def test_el_rol_global_ve_todas_las_delegaciones(cliente, crear_usuario, centro, mar):
    _vecino("Ana Centro", "11111111-1", centro)
    _vecino("Beto Mar", "12222222-5", mar)
    usuario = crear_usuario("admin@test.cl", codigo="admin", delegacion=centro)
    cliente.force_login(usuario)

    respuesta = cliente.get(reverse("ciudadanos:vecinos"))

    assert len(respuesta.context["object_list"]) == 2


def test_el_listado_de_delegaciones_se_acota_al_delegado(cliente, crear_usuario, centro, mar):
    usuario = crear_usuario("delegado@test.cl", codigo="delegado", delegacion=centro)
    cliente.force_login(usuario)

    respuesta = cliente.get(reverse("organizacion:delegaciones"))

    assert respuesta.status_code == 200
    assert [d.codigo for d in respuesta.context["object_list"]] == ["centro"]


def test_la_busqueda_es_parcial_y_no_distingue_mayusculas(cliente, crear_usuario, centro):
    _vecino("Rosa Alvarado", "11111111-1", centro)
    _vecino("Manuel Contreras", "12222222-5", centro)
    usuario = crear_usuario("admin@test.cl", codigo="admin", delegacion=centro)
    cliente.force_login(usuario)

    respuesta = cliente.get(reverse("ciudadanos:vecinos"), {"q": "ALVAR"})

    assert [v.nombre for v in respuesta.context["object_list"]] == ["Rosa Alvarado"]


def test_la_exportacion_respeta_el_filtro(cliente, crear_usuario, centro):
    _vecino("Rosa Alvarado", "11111111-1", centro)
    _vecino("Manuel Contreras", "12222222-5", centro)
    usuario = crear_usuario("admin@test.cl", codigo="admin", delegacion=centro)
    cliente.force_login(usuario)

    respuesta = cliente.get(
        reverse("ciudadanos:vecinos"), {"exportar": "1", "q": "alvar"}
    )

    assert respuesta.status_code == 200
    assert "spreadsheetml" in respuesta["Content-Type"]
    hoja = load_workbook(BytesIO(respuesta.content)).active
    filas = list(hoja.iter_rows(values_only=True))
    assert len(filas) == 2  # encabezado + la única coincidencia
    assert filas[1][0] == "Rosa Alvarado"


def test_la_ficha_de_consulta_muestra_el_registro(cliente, crear_usuario, centro):
    vecino = _vecino("Rosa Alvarado", "11111111-1", centro)
    usuario = crear_usuario("admin@test.cl", codigo="admin", delegacion=centro)
    cliente.force_login(usuario)

    respuesta = cliente.get(
        reverse("ciudadanos:vecinos_consultar", args=[vecino.pk])
    )

    assert respuesta.status_code == 200
    assert b"Rosa Alvarado" in respuesta.content
