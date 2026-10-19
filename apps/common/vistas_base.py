"""Vistas compartidas por los mantenedores.

Cada módulo declara solo sus datos: el modelo, el formulario, el título, la
plantilla y los nombres de sus rutas. El comportamiento vive acá, una vez:
búsqueda, paginación, exportación a Excel y bitácora de auditoría.
"""

from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin
from django.core.exceptions import PermissionDenied
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import redirect
from django.urls import reverse
from django.views.generic import CreateView, DeleteView, DetailView, ListView, UpdateView

from apps.common.delegaciones import acotar
from apps.common.permisos import modulo_de, puede
from apps.cumplimiento.auditoria import registrar


def resolver_campo(objeto, campo):
    """Saca el valor de un campo por ruta ("a.b") o llamando una función.

    Lo comparten la exportación a Excel y la vista de consulta: una sola forma
    de leer un campo, sin repetir el recorrido.
    """
    if callable(campo):
        return campo(objeto)
    valor = objeto
    for parte in campo.split("."):
        valor = getattr(valor, parte, "")
        if callable(valor):
            valor = valor()
    return "" if valor is None else str(valor)


class PermisoRequeridoMixin:
    """Exige el permiso del módulo según el rol, en la web.

    Va después de `LoginRequiredMixin`, así el anónimo se redirige al login y
    solo el usuario ya autenticado puede recibir un 403.
    """

    def dispatch(self, request, *args, **kwargs):
        modulo = modulo_de(getattr(self, "model", None))
        if modulo:
            escribir = request.method not in ("GET", "HEAD", "OPTIONS")
            if not puede(request.user, modulo, escritura=escribir):
                raise PermissionDenied(
                    "Tu rol no tiene permiso para este módulo."
                )
        return super().dispatch(request, *args, **kwargs)


class ExportarExcelMixin:
    """Descarga el listado en `.xlsx`.

    Reutiliza el mismo `get_queryset()` del listado, así respeta el filtro y
    exporta todas las páginas, no solo la que se está viendo.
    """

    # Cada columna es (encabezado, "ruta.atributo") o (encabezado, función).
    exportar_columnas = ()
    exportar_nombre = "listado"

    @property
    def permite_exportar(self):
        return bool(self.exportar_columnas)

    def exportar(self):
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter

        registrar(self.request, "exportar", None, f"Exportación: {self.titulo}")
        libro = Workbook()
        hoja = libro.active
        hoja.title = (self.titulo or "Listado")[:31]
        hoja.append([encabezado for encabezado, _ in self.exportar_columnas])
        for fila in self.get_queryset():
            hoja.append(
                [resolver_campo(fila, campo) for _, campo in self.exportar_columnas]
            )
        for indice, (encabezado, _) in enumerate(self.exportar_columnas, start=1):
            letra = get_column_letter(indice)
            hoja.column_dimensions[letra].width = max(14, len(encabezado) + 4)
        respuesta = HttpResponse(
            content_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        )
        respuesta["Content-Disposition"] = (
            f'attachment; filename="{self.exportar_nombre}.xlsx"'
        )
        libro.save(respuesta)
        return respuesta


class Comun:
    """Datos que cada módulo define y que las plantillas necesitan."""

    titulo = ""
    subtitulo = ""
    seccion = ""
    etiqueta_nueva = ""
    url_listado = ""
    url_nueva = None  # solo el listado la define
    permite_exportar = False

    def get_context_data(self, **kwargs):
        contexto = super().get_context_data(**kwargs)
        contexto.update(
            titulo=self.titulo,
            subtitulo=self.subtitulo,
            seccion=self.seccion,
            etiqueta_nueva=self.etiqueta_nueva,
            url_listado=self.url_listado,
            url_nueva=reverse(self.url_nueva) if self.url_nueva else "",
            permite_exportar=self.permite_exportar,
        )
        return contexto


class ListadoBase(
    ExportarExcelMixin, Comun, LoginRequiredMixin, PermisoRequeridoMixin, ListView
):
    template_name = "listado.html"
    paginate_by = 7
    busqueda = ()
    # Llaves foráneas que la plantilla recorre. Traerlas con JOIN evita una
    # consulta por fila.
    relacionadas = ()
    # Ruta del campo por el que se acota a la delegación activa. Vacío = el
    # módulo no tiene dimensión territorial (catálogos globales).
    campo_delegacion = None
    # Los listados con datos personales dejan rastro en la bitácora.
    auditar_lectura = False

    def get_queryset(self):
        queryset = super().get_queryset()
        if self.relacionadas:
            queryset = queryset.select_related(*self.relacionadas)
        q = self.request.GET.get("q")
        if q and self.busqueda:
            condicion = Q()
            for campo in self.busqueda:
                condicion |= Q(**{f"{campo}__icontains": q})
            queryset = queryset.filter(condicion)
        # El alcance territorial se aplica en el servidor, sobre el mismo
        # queryset que exporta: un cambio de contexto actualiza ambas cosas.
        if self.campo_delegacion:
            queryset = acotar(self.request, queryset, self.campo_delegacion)
        return queryset

    def get(self, request, *args, **kwargs):
        if request.GET.get("exportar") and self.exportar_columnas:
            return self.exportar()
        if self.auditar_lectura:
            registrar(request, "ver", None, f"Listado: {self.titulo}")
        return super().get(request, *args, **kwargs)


class _FormularioBase(Comun, LoginRequiredMixin, PermisoRequeridoMixin):
    template_name = "formulario.html"

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        # Los formularios acotan sus campos de delegación con la petición.
        kwargs["request"] = self.request
        return kwargs

    def get_template_names(self):
        # El modal pide el fragmento (X-Modal); sin ese encabezado se sirve la
        # página completa de siempre, que es la que ve quien no tiene JavaScript.
        if self.request.headers.get("X-Modal") == "1":
            return ["modal/formulario.html"]
        return [self.template_name]

    def get_success_url(self):
        return reverse(self.url_listado)


class AltaBase(_FormularioBase, CreateView):
    # Vacío = la vista concreta arma su propio aviso (por ejemplo, si además
    # manda un correo y quiere informar el resultado).
    mensaje_alta = "Registro creado."

    def form_valid(self, form):
        respuesta = super().form_valid(form)
        registrar(self.request, "crear", self.object)
        if self.mensaje_alta:
            messages.success(self.request, self.mensaje_alta)
        return respuesta


class EdicionBase(_FormularioBase, UpdateView):
    def form_valid(self, form):
        respuesta = super().form_valid(form)
        registrar(self.request, "editar", self.object)
        messages.success(self.request, "Cambios guardados.")
        return respuesta


class BorradoBase(Comun, LoginRequiredMixin, PermisoRequeridoMixin, DeleteView):
    template_name = "confirmar.html"

    def form_valid(self, form):
        # Se marca la fila en vez de borrarla, para no romper las referencias.
        objeto = self.object
        objeto.eliminar()
        registrar(self.request, "eliminar", objeto)
        messages.success(self.request, "Registro dado de baja.")
        return redirect(self.url_listado)


class ConsultaBase(Comun, LoginRequiredMixin, PermisoRequeridoMixin, DetailView):
    """Ficha de solo lectura de un registro.

    Completa el CRUD: el PDF pide crear, **consultar**, listar, editar y
    eliminar. Usa el mismo permiso de lectura del listado, porque no cambia
    nada. Se sirve como página completa y, con el encabezado del modal, como
    fragmento para el `<dialog>`.
    """

    template_name = "consulta.html"
    # Cada campo es (etiqueta, "ruta.atributo") o (etiqueta, función).
    consulta_campos = ()

    def get_template_names(self):
        if self.request.headers.get("X-Modal") == "1":
            return ["modal/consulta.html"]
        return [self.template_name]

    def get_context_data(self, **kwargs):
        contexto = super().get_context_data(**kwargs)
        contexto["campos"] = [
            (etiqueta, resolver_campo(self.object, campo))
            for etiqueta, campo in self.consulta_campos
        ]
        return contexto
